from pathlib import Path
from unittest.mock import Mock, patch

import pytest
from docx import Document
from pypdf import PdfWriter

from app.ai.ingestion.loaders.docx_loader import DOCXLoader
from app.ai.ingestion.loaders.pdf_loader import PDFLoader
from app.ai.ingestion.loaders.text_loader import TextLoader

# ============================================================
# TXT LOADER
# ============================================================

def test_text_loader_loads_file(tmp_path: Path) -> None:
    file_path = tmp_path / "document.txt"
    file_path.write_text(
        "Bonjour ALTIORA\nContenu du document.",
        encoding="utf-8",
    )

    loader = TextLoader()

    result = loader.load(file_path)

    assert result == "Bonjour ALTIORA\nContenu du document."


def test_text_loader_accepts_string_path(tmp_path: Path) -> None:
    file_path = tmp_path / "document.txt"
    file_path.write_text(
        "Contenu ALTIORA",
        encoding="utf-8",
    )

    loader = TextLoader()

    result = loader.load(str(file_path))

    assert result == "Contenu ALTIORA"


def test_text_loader_raises_if_file_does_not_exist(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "missing.txt"

    loader = TextLoader()

    with pytest.raises(FileNotFoundError, match="Fichier introuvable"):
        loader.load(file_path)


def test_text_loader_raises_if_source_is_directory(
    tmp_path: Path,
) -> None:
    directory = tmp_path / "documents"
    directory.mkdir()

    loader = TextLoader()

    with pytest.raises(ValueError, match="La source n'est pas un fichier"):
        loader.load(directory)


# ============================================================
# PDF LOADER
# ============================================================

def test_pdf_loader_loads_pdf(tmp_path: Path) -> None:
    file_path = tmp_path / "document.pdf"

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)

    with file_path.open("wb") as file:
        writer.write(file)

    loader = PDFLoader()

    result = loader.load(file_path)

    assert result == ""


def test_pdf_loader_accepts_string_path(tmp_path: Path) -> None:
    file_path = tmp_path / "document.pdf"

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)

    with file_path.open("wb") as file:
        writer.write(file)

    loader = PDFLoader()

    result = loader.load(str(file_path))

    assert result == ""


def test_pdf_loader_raises_if_file_does_not_exist(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "missing.pdf"

    loader = PDFLoader()

    with pytest.raises(FileNotFoundError, match="Fichier introuvable"):
        loader.load(file_path)


def test_pdf_loader_raises_if_source_is_directory(
    tmp_path: Path,
) -> None:
    directory = tmp_path / "documents"
    directory.mkdir()

    loader = PDFLoader()

    with pytest.raises(ValueError, match="La source n'est pas un fichier"):
        loader.load(directory)


def test_pdf_loader_raises_if_extension_is_invalid(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "document.txt"
    file_path.write_text("Contenu", encoding="utf-8")

    loader = PDFLoader()

    with pytest.raises(
        ValueError,
        match="La source doit être un fichier PDF",
    ):
        loader.load(file_path)

def test_pdf_loader_extracts_text_from_pages(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "document.pdf"
    file_path.write_bytes(b"%PDF-test")

    page_with_text = Mock()
    page_with_text.extract_text.return_value = "  Premier texte  "

    empty_page = Mock()
    empty_page.extract_text.return_value = "   "

    with patch(
        "app.ai.ingestion.loaders.pdf_loader.PdfReader"
    ) as pdf_reader:
        pdf_reader.return_value.pages = [
            page_with_text,
            empty_page,
        ]

        loader = PDFLoader()

        result = loader.load(file_path)

    assert result == "Premier texte"
    page_with_text.extract_text.assert_called_once()
    empty_page.extract_text.assert_called_once()


# ============================================================
# DOCX LOADER
# ============================================================

def test_docx_loader_loads_paragraphs(tmp_path: Path) -> None:
    file_path = tmp_path / "document.docx"

    document = Document()
    document.add_paragraph("Premier paragraphe")
    document.add_paragraph("Deuxième paragraphe")
    document.save(file_path)

    loader = DOCXLoader()

    result = loader.load(file_path)

    assert result == "Premier paragraphe\n\nDeuxième paragraphe"


def test_docx_loader_ignores_empty_paragraphs(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "document.docx"

    document = Document()
    document.add_paragraph("Premier paragraphe")
    document.add_paragraph("")
    document.add_paragraph("   ")
    document.add_paragraph("Deuxième paragraphe")
    document.save(file_path)

    loader = DOCXLoader()

    result = loader.load(file_path)

    assert result == "Premier paragraphe\n\nDeuxième paragraphe"


def test_docx_loader_raises_if_file_does_not_exist(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "missing.docx"

    loader = DOCXLoader()

    with pytest.raises(FileNotFoundError, match="Fichier introuvable"):
        loader.load(file_path)


def test_docx_loader_raises_if_source_is_directory(
    tmp_path: Path,
) -> None:
    directory = tmp_path / "documents"
    directory.mkdir()

    loader = DOCXLoader()

    with pytest.raises(ValueError, match="La source n'est pas un fichier"):
        loader.load(directory)


def test_docx_loader_raises_if_extension_is_invalid(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "document.txt"
    file_path.write_text("Contenu", encoding="utf-8")

    loader = DOCXLoader()

    with pytest.raises(
        ValueError,
        match="La source doit être un fichier DOCX",
    ):
        loader.load(file_path)
